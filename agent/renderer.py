# Renderer using python-docx - deterministic DOCX generation and 
# Converts Plan + SectionContent into a formatted Word doc.

import logging
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from agent.models import Plan, SectionContent

logger = logging.getLogger(__name__)


# Helpers
def _sanitize_filename(title: str) -> str:
    """Convert a document title into a filesystem-safe filename."""
    name = re.sub(r"[^\w\s-]", "", title.lower())
    name = re.sub(r"[\s]+", "_", name.strip())
    return name[:80] or "document"


def _add_table(doc: Document, table_data: list[list[str]]) -> None:
    """Insert a formatted table into the document."""
    if not table_data or len(table_data) < 2:
        return

    headers = table_data[0]
    rows = table_data[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)

    # Data rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            if idx < len(row_cells):
                row_cells[idx].text = str(value)
                for paragraph in row_cells[idx].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)


def _add_section_content(doc: Document, content_text: str) -> None:
    """Parse content text into paragraphs and bullet points."""
    lines = content_text.split("\n")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Detect bullet lines (• prefix, - prefix, or * prefix)
        if stripped.startswith(("• ", "- ", "* ")):
            bullet_text = stripped[2:].strip()
            doc.add_paragraph(bullet_text, style="List Bullet")
        else:
            doc.add_paragraph(stripped)


# Main render function
def render_docx(
    plan: Plan,
    sections: list[SectionContent],
    outputs_dir: Path,
) -> str:
    """Render a Plan + section content into a formatted .docx file.

    Args:
        plan: The document plan (title, doc_type, assumptions).
        sections: Ordered list of section content from the writer.
        outputs_dir: Directory where the .docx file will be saved.

    Returns:
        The filename of the generated document (relative to outputs_dir).
    """
    doc = Document()

    # Title
    doc.add_heading(plan.title, level=0)

    # Subtitle: document type
    subtitle = doc.add_paragraph()
    run = subtitle.add_run(f"Document Type: {plan.doc_type}")
    run.italic = True
    run.font.size = Pt(12)

    # Assumptions block (if any)
    if plan.assumptions:
        doc.add_heading("Assumptions & Notes", level=1)
        intro = doc.add_paragraph(
            "The following assumptions were made to fill gaps in the original request:"
        )
        intro.runs[0].italic = True
        for assumption in plan.assumptions:
            doc.add_paragraph(assumption, style="List Bullet")

    # Main sections
    for section in sections:
        doc.add_heading(section.heading, level=1)
        _add_section_content(doc, section.content)

        if section.table:
            _add_table(doc, section.table)

    # Save
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"{_sanitize_filename(plan.title)}_{timestamp}.docx"
    filepath = outputs_dir / filename
    doc.save(str(filepath))

    logger.info("Document rendered → %s", filepath)
    return filename
